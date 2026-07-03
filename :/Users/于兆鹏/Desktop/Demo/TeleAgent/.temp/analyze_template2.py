from pptx import Presentation
from pptx.util import Pt, Emu

path = r'D:\Training\1.培训\课程\数字化\培训\260422-23_浙江农商银行_支行长AI落地实战\1.PPT\浙江农商银行系统支行行长AI决策沙盘实战培训.pptx'
prs = Presentation(path)
slides = list(prs.slides)

print('Slide size: %.2f x %.2f inches' % (prs.slide_width/914400, prs.slide_height/914400))
print('Total slides: %d' % len(slides))

# List all layouts
print('\n=== LAYOUTS ===')
for i, layout in enumerate(prs.slide_layouts):
    phs = []
    for ph in layout.placeholders:
        phs.append('%d:%s' % (ph.placeholder_format.idx, ph.name))
    ph_str = ', '.join(phs)
    print('  Layout %d: %s | PHs: %s' % (i, layout.name, ph_str))

# Analyze all slides
print('\n=== SLIDES ===')
for i in range(len(slides)):
    slide = slides[i]
    layout = slide.slide_layout
    layout_name = layout.name if layout else 'None'
    shapes_with_text = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            if 'watermark' in shape.name.lower():
                continue
            txt = shape.text.strip()[:50].replace('\n', ' | ')
            shapes_with_text.append('%s=%s' % (shape.name, txt))
    shapes_brief = ' | '.join(shapes_with_text[:4]) if shapes_with_text else '(no text)'
    print('  Slide%d [%s]: %s' % (i+1, layout_name, shapes_brief[:180]))

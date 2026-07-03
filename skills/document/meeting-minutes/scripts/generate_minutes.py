#!/usr/bin/env python3
"""Generate a formatted DOCX meeting minutes document from structured JSON data.

Usage:
    python generate_minutes.py --input minutes_data.json --output meeting_minutes.docx

Input JSON format:
{
    "title": "会议名称",
    "time": "2025-01-15 14:00-16:00",
    "location": "会议室A301",
    "participants": [
        {"name": "张三", "role": "主持人"},
        {"name": "李四", "role": "参会人"}
    ],
    "absent": ["王五"],
    "decisions": [
        {"id": 1, "content": "决议内容", "notes": "补充说明"}
    ],
    "action_items": [
        {"id": 1, "task": "待办任务描述", "assignee": "负责人", "deadline": "截止日期", "status": "待启动"}
    ],
    "notes": "其他备注信息"
}
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    print("ERROR: python-docx is required. Install with: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear'
    })
    shading.append(shading_elem)


def set_run_font(run, font_name_cn='微软雅黑', font_name_en='Calibri', size=Pt(10.5),
                 bold=False, color=None):
    """Configure font for a run."""
    run.font.size = size
    run.font.bold = bold
    run.font.name = font_name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name_cn)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_styled(doc, text, level=1):
    """Add a styled heading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return heading


def build_minutes(data):
    """Build a Document from structured meeting minutes data."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # Title
    title_text = data.get('title', '会议纪要')
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title_text)
    set_run_font(title_run, size=Pt(22), bold=True, color=(0x1F, 0x49, 0x7D))

    # Separator line
    sep_para = doc.add_paragraph()
    sep_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sep_run = sep_para.add_run('━' * 40)
    set_run_font(sep_run, size=Pt(10), color=(0xBF, 0xBF, 0xBF))

    # Basic info table
    add_heading_styled(doc, '基本信息', level=2)

    info_table = doc.add_table(rows=3, cols=4, style='Table Grid')
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    info_items = [
        ('会议时间', data.get('time', '未指定')),
        ('会议地点', data.get('location', '未指定')),
    ]

    # Row 0: time
    cell = info_table.cell(0, 0)
    cell.text = ''
    run = cell.paragraphs[0].add_run('会议时间')
    set_run_font(run, bold=True, size=Pt(10.5))
    set_cell_shading(cell, 'E8F0FE')

    cell = info_table.cell(0, 1)
    cell.text = ''
    run = cell.paragraphs[0].add_run(data.get('time', '未指定'))
    set_run_font(run, size=Pt(10.5))
    # Merge remaining cells for time
    cell_merge = info_table.cell(0, 1)
    for c in [info_table.cell(0, 2), info_table.cell(0, 3)]:
        cell_merge = cell_merge.merge(c)
    cell_merge.text = ''
    run = cell_merge.paragraphs[0].add_run(data.get('time', '未指定'))
    set_run_font(run, size=Pt(10.5))

    # Row 1: location
    cell = info_table.cell(1, 0)
    cell.text = ''
    run = cell.paragraphs[0].add_run('会议地点')
    set_run_font(run, bold=True, size=Pt(10.5))
    set_cell_shading(cell, 'E8F0FE')

    cell_merge = info_table.cell(1, 1)
    for c in [info_table.cell(1, 2), info_table.cell(1, 3)]:
        cell_merge = cell_merge.merge(c)
    cell_merge.text = ''
    run = cell_merge.paragraphs[0].add_run(data.get('location', '未指定'))
    set_run_font(run, size=Pt(10.5))

    # Row 2: participants
    cell = info_table.cell(2, 0)
    cell.text = ''
    run = cell.paragraphs[0].add_run('参会人员')
    set_run_font(run, bold=True, size=Pt(10.5))
    set_cell_shading(cell, 'E8F0FE')

    participants = data.get('participants', [])
    if participants:
        p_text = '、'.join(
            f"{p['name']}({p['role']})" if p.get('role') else p['name']
            for p in participants
        )
    else:
        p_text = '未指定'

    cell_merge = info_table.cell(2, 1)
    for c in [info_table.cell(2, 2), info_table.cell(2, 3)]:
        cell_merge = cell_merge.merge(c)
    cell_merge.text = ''
    run = cell_merge.paragraphs[0].add_run(p_text)
    set_run_font(run, size=Pt(10.5))

    # Absent
    absent = data.get('absent', [])
    if absent:
        doc.add_paragraph()
        absent_para = doc.add_paragraph()
        absent_run = absent_para.add_run('缺席人员：')
        set_run_font(absent_run, bold=True, size=Pt(10.5))
        absent_run2 = absent_para.add_run('、'.join(absent))
        set_run_font(absent_run2, size=Pt(10.5))

    # Decisions
    decisions = data.get('decisions', [])
    if decisions:
        add_heading_styled(doc, '决议事项', level=2)
        for d in decisions:
            para = doc.add_paragraph()
            num_run = para.add_run(f"{d.get('id', '')}. " if d.get('id') else '- ')
            set_run_font(num_run, bold=True, size=Pt(10.5), color=(0x1F, 0x49, 0x7D))
            content_run = para.add_run(d.get('content', ''))
            set_run_font(content_run, size=Pt(10.5))
            if d.get('notes'):
                notes_para = doc.add_paragraph()
                notes_para.paragraph_format.left_indent = Cm(0.75)
                notes_run = notes_para.add_run(f"补充：{d['notes']}")
                set_run_font(notes_run, size=Pt(9), color=(0x66, 0x66, 0x66))

    # Action items
    action_items = data.get('action_items', [])
    if action_items:
        add_heading_styled(doc, '待办任务', level=2)

        ai_table = doc.add_table(rows=1 + len(action_items), cols=5, style='Table Grid')
        ai_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        headers = ['序号', '任务', '负责人', '截止日期', '状态']
        for i, h in enumerate(headers):
            cell = ai_table.cell(0, i)
            cell.text = ''
            run = cell.paragraphs[0].add_run(h)
            set_run_font(run, bold=True, size=Pt(10.5))
            set_cell_shading(cell, '1F497D')
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Rows
        for idx, item in enumerate(action_items):
            row = idx + 1
            values = [
                str(item.get('id', idx + 1)),
                item.get('task', ''),
                item.get('assignee', ''),
                item.get('deadline', ''),
                item.get('status', '待启动'),
            ]
            for col, val in enumerate(values):
                cell = ai_table.cell(row, col)
                cell.text = ''
                run = cell.paragraphs[0].add_run(val)
                set_run_font(run, size=Pt(10))

    # Notes
    notes = data.get('notes', '')
    if notes:
        add_heading_styled(doc, '其他备注', level=2)
        notes_para = doc.add_paragraph()
        notes_run = notes_para.add_run(notes)
        set_run_font(notes_run, size=Pt(10.5))

    return doc


def main():
    parser = argparse.ArgumentParser(description='Generate meeting minutes DOCX')
    parser.add_argument('--input', required=True, help='Input JSON file path')
    parser.add_argument('--output', required=True, help='Output DOCX file path')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: Input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    with open(input_path, 'r', encoding='utf-8-sig') as f:
        data = json.load(f)

    doc = build_minutes(data)
    doc.save(args.output)
    print(f"Meeting minutes saved to: {args.output}")


if __name__ == '__main__':
    main()
